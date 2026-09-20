from pathlib import Path

import pymupdf
from docx import Document

from regagent.infrastructure.parsers import ParserRegistry
from regagent.infrastructure.parsers.html import HtmlParser


def test_html_parser_preserves_visible_block_order() -> None:
    path = Path(__file__).parents[1] / "fixtures" / "sample_regulation_ru.html"

    parsed = ParserRegistry().parse(path)

    assert parsed.title == "Тестовый нормативный акт"
    assert parsed.blocks[0].text == "Тестовый нормативный акт"
    assert parsed.blocks[3].text.startswith("Статья 1")


def test_docx_parser_reads_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_heading("Нормативный документ", level=1)
    document.add_paragraph("Статья 1. Общие положения")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Срок"
    table.cell(0, 1).text = "30 дней"
    document.save(path)

    parsed = ParserRegistry().parse(path)

    assert parsed.title == "Нормативный документ"
    assert any(block.text == "Срок | 30 дней" for block in parsed.blocks)


def test_pdf_parser_reads_positioned_text_blocks(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Article 1. Test provision")
    document.save(path)
    document.close()

    parsed = ParserRegistry().parse(path)

    assert parsed.media_type == "application/pdf"
    assert parsed.blocks[0].page == 1
    assert "Article 1" in parsed.blocks[0].text


def test_html_parser_splits_explicit_break_between_legal_headings(tmp_path: Path) -> None:
    path = tmp_path / "headings.html"
    path.write_text(
        "<html><body><main><h3>РАЗДЕЛ 1. ОБЩИЕ ПОЛОЖЕНИЯ"
        "<br>Глава 1. ОСНОВНЫЕ ПОЛОЖЕНИЯ</h3></main></body></html>",
        encoding="utf-8",
    )

    parsed = HtmlParser().parse(path)

    assert [block.text for block in parsed.blocks] == [
        "РАЗДЕЛ 1. ОБЩИЕ ПОЛОЖЕНИЯ",
        "Глава 1. ОСНОВНЫЕ ПОЛОЖЕНИЯ",
    ]
